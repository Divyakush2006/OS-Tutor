"""Turn a course document into Markdown whose headings are the citation anchors.

Supported: .pdf, .pptx, .docx, .md, .txt. Legacy .ppt/.doc are converted first with
tools/convert_office_legacy.ps1 (the installed Office apps do it losslessly).

  .pptx  "## Slide N: <slide title>" per slide; the body text of every shape (groups and
         tables included), then the speaker notes.
  .pdf   "## Page N" per page; lines set in a clearly larger font than the page body are
         promoted to "### <heading>", so a citation reads "Page 3 › 28.1 Locks: The Basic Idea".
         Hyphenated line breaks are joined; bare page numbers are dropped.
  .docx  Word "Title" / "Heading N" styles become headings; tables become "a | b | c" rows.
  .md    used as is; .txt gets no headings (citations then name the document only).
"""
from __future__ import annotations

import re
import statistics
import unicodedata
from pathlib import Path

SUPPORTED = {".pdf", ".pptx", ".docx", ".md", ".txt"}


def extract(path: Path, title: str) -> str:
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"{path.name}: unsupported format {ext} (supported: {sorted(SUPPORTED)})")
    body = {".pdf": _pdf, ".pptx": _pptx, ".docx": _docx}.get(ext, _text)(path)
    # NFKC folds typographic ligatures ("ﬁle" -> "file") and full-width forms, so a
    # student's plain-ASCII query matches the text of a typeset PDF
    body = unicodedata.normalize("NFKC", body)
    if ext != ".md" or not body.lstrip().startswith("# "):
        body = f"# {title}\n\n{body}"
    return re.sub(r"\n{3,}", "\n\n", body).strip() + "\n"


def _text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


# ---------------------------------------------------------------- PowerPoint

def _shape_text(shape) -> list[str]:
    out = []
    if shape.shape_type == 6:                                   # group: recurse
        for s in shape.shapes:
            out += _shape_text(s)
    elif getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs).strip()
            if t:
                out.append(("- " if p.level else "") + t)
    return out


def _pptx(path: Path) -> str:
    from pptx import Presentation
    deck = Presentation(str(path))
    parts = []
    for n, slide in enumerate(deck.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip().replace("\n", " ") if title_shape is not None else ""
        lines = []
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            lines += _shape_text(shape)
        lines = [l for l in lines if not re.fullmatch(r"\d{1,3}", l)]      # slide numbers
        if not title and lines:
            title = lines.pop(0)[:80]
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        parts.append(f"## Slide {n}: {title or '(untitled)'}\n\n" + "\n".join(lines)
                     + (f"\n\nNotes: {notes}" if notes else ""))
    return "\n\n".join(parts)


# ---------------------------------------------------------------- PDF

def _running(line: str) -> str:
    """A line with its digits removed, so page headers that carry a page number compare equal."""
    return re.sub(r"\d+", "", line).strip().lower()


def _pdf(path: Path) -> str:
    import fitz
    doc = fitz.open(str(path))
    pages = []
    for page in doc:
        lines = []                                              # (size, text)
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                spans = [s for s in line["spans"] if s["text"].strip()]
                if spans:
                    size = max(s["size"] for s in spans)
                    lines.append((size, "".join(s["text"] for s in line["spans"]).strip()))
        pages.append(lines)
    # Running headers and footers ("OPERATING SYSTEMS [VERSION 1.01] WWW.OSTEP.ORG", the
    # chapter name in capitals, a copyright line) repeat on many pages and would be glued
    # into the middle of passages. A line on at least a third of the pages (min. 3) is one.
    seen = {}
    for lines in pages:
        for key in {_running(t) for _, t in lines}:
            seen[key] = seen.get(key, 0) + 1
    limit = max(3, len(pages) / 3)
    running = {k for k, c in seen.items() if c >= limit and k}
    parts = []
    for n, lines in enumerate(pages, 1):
        lines = [(s, t) for s, t in lines if _running(t) not in running]
        if not lines:
            continue
        sizes = [s for s, t in lines for _ in range(len(t))]   # weight by characters
        body = statistics.median(sizes) if sizes else 0
        out = [f"## Page {n}"]
        for size, text in lines:
            if re.fullmatch(r"\d{1,4}", text):                  # page numbers
                continue
            if size >= body * 1.25 and len(text) <= 90 and re.search(r"[A-Za-z]{3}", text):
                out.append(f"\n### {text}\n")
            else:
                out.append(text)
        page_md = "\n".join(out)
        page_md = re.sub(r"(\w)-\n(\w)", r"\1\2", page_md)      # join hyphenated breaks
        parts.append(page_md)
    return "\n\n".join(parts)


# ---------------------------------------------------------------- Word

def _docx(path: Path) -> str:
    import docx
    d = docx.Document(str(path))
    out = []
    body = d.element.body
    paras = {p._p: p for p in d.paragraphs}
    tables = {t._tbl: t for t in d.tables}
    for el in body.iterchildren():
        if el in paras:
            p = paras[el]
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            m = re.match(r"heading (\d)", style)
            if style == "title":
                out.append(f"\n## {t}\n")
            elif m:
                out.append(f"\n{'#' * min(int(m.group(1)) + 1, 4)} {t}\n")
            else:
                out.append(t)
        elif el in tables:
            for row in tables[el].rows:
                out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)
