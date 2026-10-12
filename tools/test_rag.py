"""Tests for the RAG pipeline (G2). Hermetic: no GPU, no model server, no network.

    python tools/test_rag.py
"""
from __future__ import annotations

import json
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from ostutor.rag.corpus import MAX_WORDS, CorpusError, load_corpus, sha256  # noqa: E402
from ostutor.rag.extract import extract  # noqa: E402
from ostutor.rag.ingest import add  # noqa: E402
from ostutor.rag.retrieve import (BM25, Retriever, augment, cited,  # noqa: E402
                                  sources_footer, tokens)

FIXTURE = ROOT / "data" / "rag" / "fixture"
CORPUS = ROOT / "data" / "rag" / "corpus"
FAILED: list[str] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    print(f"  {'PASS' if ok else 'FAIL'}  {name}" + (f"  -> {detail}" if not ok and detail else ""))
    if not ok:
        FAILED.append(name)


def corpus(docs: dict[str, str], meta: dict | None = None) -> Path:
    """A throwaway corpus with a complete manifest; `meta` overrides fields."""
    d = Path(tempfile.mkdtemp())
    entries = []
    for i, (name, text) in enumerate(docs.items()):
        (d / name).write_text(text, encoding="utf-8")
        entries.append({**dict(id=f"d{i}", title=name, path=name, source="test", owner="tester",
                               licence="CC0-1.0", permission="test", added="2026-10-12",
                               sha256=sha256(d / name)), **(meta or {})})
    (d / "manifest.json").write_text(json.dumps(dict(documents=entries)), encoding="utf-8")
    return d


def refused(d: Path) -> bool:
    try:
        load_corpus(d)
        return False
    except CorpusError:
        return True


print("corpus: manifest rules")
chunks, man = load_corpus(FIXTURE)
check("the fixture loads", len(chunks) > 0 and man["chunk_count"] == len(chunks))
check("every chunk carries its note and section", all(c.doc_id and c.section for c in chunks))
check("every chunk carries a licence", all(c.licence for c in chunks))
check("no chunk exceeds the size limit", max(len(c.text.split()) for c in chunks) <= MAX_WORDS)
check("chunk ids are unique", len({c.chunk_id for c in chunks}) == len(chunks))
check("the corpus hash is stable", load_corpus(FIXTURE)[1]["chunks_sha256"] == man["chunks_sha256"])

for field in ("source", "owner", "added", "sha256"):
    d = corpus({"a.md": "## S\nsome words here for a chunk"}, meta={field: ""})
    check(f"a document with no {field} is refused", refused(d))
    shutil.rmtree(d)
d = corpus({"a.md": "## S\nsome words here for a chunk"}, meta={"licence": "", "permission": ""})
check("a document with neither licence nor permission is refused", refused(d))
shutil.rmtree(d)
d = corpus({"a.md": "## S\nsome words here for a chunk"}, meta={"licence": ""})
check("a permission alone is enough", not refused(d))
shutil.rmtree(d)

d = corpus({"a.md": "## S\nsome words here for a chunk"})
(d / "a.md").write_text("## S\ntampered words here for a chunk", encoding="utf-8")
check("a file changed after it was recorded is refused (checksum)", refused(d))
shutil.rmtree(d)

d = corpus({"a.md": "## S\nsome words here for a chunk"})
(d / "stray.pdf").write_bytes(b"%PDF-1.4")
check("a file with no manifest entry is refused", refused(d))
(d / "stray.pdf").unlink()
(d / "extracted").mkdir()
(d / "extracted" / "x.md").write_text("cache", encoding="utf-8")
check("derived files under extracted/ are not documents", not refused(d))
shutil.rmtree(d)

d = corpus({"a.md": "# Doc\n## One\nalpha beta gamma delta epsilon zeta\n### Inner\neta theta iota kappa lambda mu\n"
                    "## Two\nnu xi omicron pi rho sigma\n"})
c, _ = load_corpus(d)
check("headings nest into a section path",
      [x.section for x in c] == ["One", "One › Inner", "Two"], str([x.section for x in c]))
shutil.rmtree(d)

d = corpus({"long.md": "## S\n" + " ".join(f"w{i}" for i in range(400))})
c, _ = load_corpus(d)
check("a long section is split with overlap",
      len(c) == 3 and c[0].text.split()[-30:] == c[1].text.split()[:30], str(len(c)))
shutil.rmtree(d)

print("ingest")
d = Path(tempfile.mkdtemp())
(d / "docs").mkdir()
(d / "docs" / "w.md").write_text("## Week 1\nprocesses and threads and scheduling basics", encoding="utf-8")
add(d, [str(d / "docs" / "w.md")], owner="Instructor", source="LMS", licence="", permission="course use")
e = json.loads((d / "manifest.json").read_text(encoding="utf-8"))["documents"]
check("add records owner, source, date and checksum",
      len(e) == 1 and e[0]["owner"] == "Instructor" and e[0]["added"] and e[0]["sha256"] == sha256(d / "docs" / "w.md"))
check("an added file then ingests", not refused(d))
try:
    add(d, [str(d / "docs" / "w.md")], owner="x", source="y", licence="", permission="")
    check("add refuses a file with neither licence nor permission", False)
except SystemExit:
    check("add refuses a file with neither licence nor permission", True)
shutil.rmtree(d)

print("extraction")
tmp = Path(tempfile.mkdtemp())
from pptx import Presentation  # noqa: E402
deck = Presentation()
s = deck.slides.add_slide(deck.slide_layouts[1])
s.shapes.title.text = "Belady's Anomaly"
s.placeholders[1].text = "FIFO can fault more with more frames"
deck.save(tmp / "d.pptx")
md = extract(tmp / "d.pptx", "Deck")
check("a slide's title becomes its section", "## Slide 1: Belady's Anomaly" in md, md[:120])
check("slide body text is kept", "FIFO can fault more with more frames" in md)

import docx  # noqa: E402
w = docx.Document()
w.add_heading("Deadlock", level=1)
w.add_paragraph("Four conditions must hold.")
t = w.add_table(rows=1, cols=2)
t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Mutual exclusion", "Hold and wait"
w.save(tmp / "d.docx")
md = extract(tmp / "d.docx", "Handout")
check("a Word heading becomes a section", "## Deadlock" in md, md[:120])
check("Word tables are kept as rows", "Mutual exclusion | Hold and wait" in md)

import fitz  # noqa: E402
pdf = fitz.open()
for n in range(4):
    p = pdf.new_page()
    p.insert_text((72, 40), "OPERATING SYSTEMS [VERSION 1.01] WWW.EXAMPLE.ORG", fontsize=9)
    p.insert_text((72, 90), f"{n + 1}.1 Topic {'ABCD'[n]} Section", fontsize=18)
    p.insert_text((72, 130), f"Body text number {n} about lock {'wxyz'[n]} and mutual exclusion.", fontsize=10)
pdf.save(tmp / "d.pdf")
md = extract(tmp / "d.pdf", "Chapter")
check("each PDF page is a section", "## Page 1" in md and "## Page 4" in md, md[:200])
check("large-font lines become headings", "### 1.1 Topic A Section" in md, md[:200])
check("running headers repeated on every page are dropped", "WWW.EXAMPLE.ORG" not in md)
(tmp / "lig.md").write_text("## S\nThe ﬁrst thread acquires the ﬂag.", encoding="utf-8")
check("ligatures are folded so 'first' matches", "The first thread acquires the flag" in extract(tmp / "lig.md", "L"))
shutil.rmtree(tmp)

print("bm25")
bm = BM25([tokens("page replacement belady anomaly fifo"), tokens("disk scheduling elevator scan")])
check("the matching document scores highest", bm.scores(tokens("belady fifo")).argmax() == 0)
check("an unrelated query scores zero", bm.scores(tokens("javascript closure")).max() == 0)
check("stop words are ignored", tokens("What is the TLB?") == ["tlb"])

print("retrieval (BM25 only)")
r = Retriever(FIXTURE, dense=False)
hits = r.retrieve("Can LRU suffer from Belady's anomaly?", mode="bm25")
check("finds the Belady section", hits and hits[0]["chunk"].section == "Belady's anomaly")
check("returns at most k chunks", len(r.retrieve("page", k=2, mode="bm25")) <= 2)
check("a query sharing no word with the notes returns nothing",
      r.retrieve("javascript closure", mode="bm25") == [])

print("retrieval (hybrid, local embedding model)")
rh = Retriever(FIXTURE)
h = rh.retrieve("The system spends nearly all its time swapping pages and the CPU is idle")
check("a paraphrase finds thrashing", h and h[0]["chunk"].doc_id == "n08", str(h[:1]))
check("an off-topic question abstains", rh.retrieve("What is the capital of Australia?") == [])
check("gate='none' never abstains", len(rh.retrieve("What is the capital of Australia?", gate="none")) > 0)
if (CORPUS / "gate.json").is_file():
    g = json.loads((CORPUS / "gate.json").read_text(encoding="utf-8"))
    rc = Retriever(CORPUS, dense=False)
    check("the course corpus loads with its calibrated gate",
          len(rc.chunks) > 1000 and rc.strong_cosine == g["strong_cosine"] and rc.min_bm25 == g["min_bm25"])
    kinds = {d["kind"] for d in rc.manifest["documents"]}
    check("every course-corpus document says whether it is course or online material",
          kinds == {"course", "online"}, str(kinds))

print("prompt and citations")
SYSTEM = "You are an Operating Systems tutor."
check("no hits leaves the trained prompt untouched", augment(SYSTEM, []) == SYSTEM)
aug = augment(SYSTEM, hits[:2])
check("the trained prompt is kept word for word at the start", aug.startswith(SYSTEM + "\n\n"))
check("each passage is labelled [S1], [S2]", "[S1] Page replacement" in aug and "[S2]" in aug)
check("the model is told never to invent a citation", "Never invent a citation" in aug)
check("the footer names each source", "`n07`" in sources_footer(hits[:2]))
check("no hits, no footer", sources_footer([]) == "")
check("valid and invented citations are told apart",
      cited("as [S1] and [S2] show, and [S7]", 2) == ({1, 2}, {7}))

print("\nALL RAG TESTS PASSED" if not FAILED else f"\n{len(FAILED)} FAILED: {FAILED}")
sys.exit(1 if FAILED else 0)
